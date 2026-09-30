# -*- coding: utf-8 -*-
"""阶段3：生成 config_table.xlsx 的「翻译」与「术语表」sheet（初版回填）。

四股来源并集，去重：
  1. config/loc_glossary.json 全部条目
  2. config/*.json 展示字段值：buff_name/buff_tooltip/debuff_name/debuff_tooltip/
     coin_skill_name/coin_skill_tooltip/card_name
     （⚠️ 排除 card_description —— 运行时被 _generate_card_description() 覆盖）
  3. scripts/**/*.gd 中 tr("...") / T("...") 的常量参数
  4. scenes/**/*.tscn 中 text / tooltip_text / popup/item_<n>/text 属性的中文值

「术语表」sheet 收录：glossary 中"纯文本专名类"条目（不含占位符/BBCode/正则替换标记）。
正式运行前自动备份 xlsx；支持 --dry 试运行。

用法：
  py -3.14 -X utf8 tools/gen_initial_sheet.py --dry
  py -3.14 -X utf8 tools/gen_initial_sheet.py
运行前请关闭 Excel。xlsx 结构无损：只新增两个 sheet，不动其余 sheet 数据区。
"""
import json
import os
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "config" / "config_table.xlsx"
LOCALES = ["en", "ja", "ko", "zh_TW", "fr", "de", "es", "pt_BR", "it"]
DRY = "--dry" in sys.argv

JSON_DISPLAY_FIELDS = [
    "buff_name", "buff_tooltip",
    "debuff_name", "debuff_tooltip",
    "coin_skill_name", "coin_skill_tooltip",
    "card_name",
]

# 含汉字（CJK 统一表意文字区间）
HAN = re.compile(r"[\u4e00-\u9fff\u3400-\u4dbf]")

# tscn 中 text / tooltip_text / OptionButton item text 的属性值（含跨行与转义引号）
TSCN_PROP = re.compile(r"(?m)^(?:popup/item_\d+/)?(?:text|tooltip_text)\s*=\s*\"((?:\\.|[^\"\\])*)\"")
# gd 中 tr( / T( 的括号配对提取：tr("literal")
GD_CALL = re.compile(r"\b(?:tr|T)\s*\(\s*\"((?:\\.|[^\"\\])*)\"")

# glossary 中属于"转义/模板工具串"而非专名的标记：占位符、BBCode、$组引用
NOT_GLOSSARY_TERM = re.compile(r"%[sdf]|\[/?[a-z]|^\s*$|\$")

# 已废弃/可结束空串等噪音
def is_noise(s: str) -> bool:
    if not s or not s.strip():
        return True
    return HAN.search(s) is None and not NOT_GLOSSARY_TERM.search(s)


def unescape(s: str) -> str:
    """还原 GDScript/TSCN 字符串转义"""
    return (s.replace("\\r\\n", "\n").replace("\\n", "\n")
             .replace("\\t", "\t").replace("\\r", "\n")
             .replace("\\\"", "\"").replace("\\\\", "\\"))


def collect_glossary():
    items = json.load(open(ROOT / "config" / "loc_glossary.json", encoding="utf-8"))
    return [str(x) for x in items]


def collect_json_fields():
    import glob as _glob
    vals = []
    for fp in sorted(_glob.glob(str(ROOT / "config" / "*.json"))):
        if Path(fp).name == "loc_glossary.json":
            continue
        data = json.load(open(fp, encoding="utf-8"))
        if not isinstance(data, list):
            continue
        for row in data:
            for field in JSON_DISPLAY_FIELDS:
                v = row.get(field)
                if isinstance(v, str) and v.strip():
                    vals.append(v)
    return vals


def collect_gd_literals():
    vals = []
    for fp in sorted((ROOT / "scripts").rglob("*.gd")):
        code = fp.read_text(encoding="utf-8")
        for m in GD_CALL.finditer(code):
            lit = unescape(m.group(1))
            if lit.strip() and lit != "，":  # 纯","分隔符单独登记无意义，跳过
                vals.append(lit)
    return vals


def collect_tscn_values():
    vals = []
    for fp in sorted((ROOT / "scenes").rglob("*.tscn")):
        content = fp.read_text(encoding="utf-8")
        for m in TSCN_PROP.finditer(content):
            lit = unescape(m.group(1)).strip()
            if lit and HAN.search(lit):
                vals.append(lit)
    return vals


def main():
    src_glossary = collect_glossary()
    src_json = collect_json_fields()
    src_gd = collect_gd_literals()
    src_tscn = collect_tscn_values()

    keys = []                 # 有序去重
    seen = set()

    def add(s):
        if s not in seen:
            seen.add(s)
            keys.append(s)

    for s in src_glossary:
        add(s)
    for s in src_json:
        add(s)
    for s in src_tscn:
        add(s)
    for s in src_gd:
        add(s)

    glossary_terms = [s for s in src_glossary if not NOT_GLOSSARY_TERM.search(s)]

    print("==== 来源统计 ====")
    print(f"  1. loc_glossary.json      : {len(src_glossary)} 条")
    print(f"  2. config/*.json 展示字段 : {len(src_json)} 条")
    print(f"  3. scripts tr()/T() 字面量: {len(src_gd)} 条")
    print(f"  4. scenes text/tooltip    : {len(src_tscn)} 条")
    print(f"  并集去重后（翻译 sheet 行数）: {len(keys)} 条")
    print(f"  术语表专名条目（来自 glossary 纯文本）: {len(glossary_terms)} 条")

    noise = [s for s in keys if is_noise(s)]
    if noise:
        print(f"  ⚠ 疑似噪音（无汉字且不含占位符，保留供人工确认）：{len(noise)} 条")
        for s in noise[:10]:
            print("   ", repr(s))

    # 同形多义风险清单：短 key 且同时出现在 场景/代码 与 数据 的
    crossed = set(src_json) & (set(src_gd) | set(src_tscn))
    if crossed:
        print(f"  ⚠ 同文多义排查（数据与代码/场景同形）：{len(crossed)} 条")
        for s in sorted(crossed):
            print("   ", repr(s))

    if DRY:
        print("\n（试运行结束，未写盘。去掉 --dry 正式执行）")
        return

    bak = XLSX.with_name(f"config_table.backup-{datetime.now():%Y%m%d-%H%M%S}.xlsx")
    shutil.copy2(XLSX, bak)
    print(f"\n已备份: {bak.name}")

    wb = load_workbook(XLSX)
    for name in ("翻译", "术语表"):
        if name in wb.sheetnames:
            del wb[name]
    ws = wb.create_sheet("翻译")
    ws.cell(row=1, column=1, value="keys")
    for ci, loc in enumerate(LOCALES, start=2):
        ws.cell(row=1, column=ci, value=loc)
    ws.cell(row=1, column=len(LOCALES) + 2, value="备注")
    for ri, key in enumerate(keys, start=2):
        ws.cell(row=ri, column=1, value=key)

    ws2 = wb.create_sheet("术语表")
    ws2.cell(row=1, column=1, value="中文")
    for ci, loc in enumerate(LOCALES, start=2):
        ws2.cell(row=1, column=ci, value=loc)
    ws2.cell(row=1, column=len(LOCALES) + 2, value="备注")
    for ri, key in enumerate(glossary_terms, start=2):
        ws2.cell(row=ri, column=1, value=key)

    try:
        wb.save(XLSX)
    except PermissionError:
        print("保存失败：config_table.xlsx 正被 Excel 占用，请关闭 Excel 后重跑（备份未动）")
        sys.exit(1)
    print(f"已写入「翻译」sheet：{len(keys)} 行；「术语表」sheet：{len(glossary_terms)} 行")

    # ---- 回读校验 ----
    wb2 = load_workbook(XLSX, read_only=True)
    assert "翻译" in wb2.sheetnames and "术语表" in wb2.sheetnames
    rows = list(wb2["翻译"].iter_rows(values_only=True))
    header = rows[0]
    assert list(header[:2]) == ["keys", "en"], f"表头不符: {header[:2]}"
    body = [r[0] for r in rows[1:] if r[0]]
    assert body == keys, "回读与写入不一致"
    g2 = list(wb2["术语表"].iter_rows(values_only=True))
    gbody = [r[0] for r in g2[1:] if r[0]]
    assert gbody == glossary_terms, "术语表回读与写入不一致"
    print("回读校验 ✓ 翻译/术语表内容与预期一致")


if __name__ == "__main__":
    main()
